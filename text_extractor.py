"""Resume text extraction (PDF via PyMuPDF, DOCX via python-docx)."""
import pymupdf as fitz  # PyMuPDF
from docx import Document


class ExtractionError(Exception):
    pass


def _extract_pdf(path: str) -> tuple[str, dict]:
    try:
        doc = fitz.open(path)
    except Exception as exc:
        raise ExtractionError(f"Could not open PDF: {exc}")
    parts, images = [], 0
    for page in doc:
        parts.append(page.get_text("text"))
        images += len(page.get_images())
    meta = {"pages": doc.page_count, "images": images, "tables": 0}
    doc.close()
    return "\n".join(parts), meta


def _extract_docx(path: str) -> tuple[str, dict]:
    try:
        doc = Document(path)
    except Exception as exc:
        raise ExtractionError(f"Could not open DOCX: {exc}")
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in cells:
                    cells.append(t)
            if cells:
                lines.append(" | ".join(cells))
    text = "\n".join(lines)
    words = len(text.split())
    meta = {"pages": max(1, round(words / 500)), "images": len(doc.inline_shapes), "tables": len(doc.tables)}
    return text, meta


def extract_text(path: str, file_type: str) -> tuple[str, dict]:
    file_type = file_type.lower().lstrip(".")
    if file_type == "pdf":
        text, meta = _extract_pdf(path)
    elif file_type == "docx":
        text, meta = _extract_docx(path)
    else:
        raise ExtractionError("Unsupported file type. Upload a PDF or DOCX file.")
    meta["word_count"] = len(text.split())
    return text, meta
